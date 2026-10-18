"""Read the supplied workbook without changing it; emit the mod's vendor catalogue."""
import openpyxl,json,re,sys,pathlib
world_names=['Spawn','Nether','End','Heaven','Ocean','Acidic Lakes','Steampunk','Sakura Forest','Wild West','Pixel Grove','Planet X']
w=openpyxl.load_workbook(sys.argv[1],data_only=True)
worlds=[]
for name in world_names:
    mines=[]
    if name in w.sheetnames:
        sheet=w[name]
        starts=[r for r in range(1,sheet.max_row+1) if re.match(r'MINE\s+\d+',str(sheet.cell(r,3).value or ''),re.I)]
        for start in starts:
            label=re.sub(r'\s+',' ',sheet.cell(start,3).value.strip())
            resource=str(sheet.cell(start+2,3).value or '').partition(':')[2].strip()
            header=start+4
            tool=str(sheet.cell(header,4).value or 'Tool')
            rows=[]; notes=[]
            for offset in range(5):
                r=header+1+offset
                values=[]
                for col in (3,4,5):
                    v=sheet.cell(r,col).value
                    values.append('Not recorded' if v is None else str(v).replace(' ',''))
                rows.append({'tier':str(sheet.cell(r,1).value),'sword':values[0],'tool':values[1],'armor':values[2]})
            specials={}
            for key,col in [('talisman',6),('elytra',7)]:
                cell=sheet.cell(header+1,col); v=cell.value
                if v is None: value='Not recorded'
                elif isinstance(v,(int,float)):
                    value=str(v)+' (unit unconfirmed)'
                    notes.append(f'{cell.coordinate}: unit missing in source; verify with vendor.')
                elif not re.fullmatch(r'[\d.,]+\s*[KMBTkmbt]?',str(v).strip()):
                    value='Not confirmed'
                    notes.append(f'{cell.coordinate}: source says "{v}"; verify with vendor.')
                else: value=str(v).replace(' ','')
                specials[key]=value
            mines.append({'name':label,'resource':resource,'toolName':tool,'rows':rows,**specials,'notes':notes,'source':f'Techmines.xlsx / {name} / C{start}:G{header+5}'})
    worlds.append({'name':name,'mines':mines})
output={'schemaVersion':1,'source':'User supplied Techmines.xlsx; imported 2026-10-18','worlds':worlds}
p=pathlib.Path(sys.argv[2]);p.write_text(json.dumps(output,indent=2),encoding='utf-8')
assert len(worlds)==11
assert sum(len(x['mines']) for x in worlds)==16
assert worlds[9]['mines'][0]['rows'][0]['sword']=='222.35M'
assert 'unit unconfirmed' in worlds[9]['mines'][1]['talisman']
assert worlds[10]['mines'][7]['elytra']=='Not confirmed'
print('Imported 11 worlds, 16 mines; preserved missing costs and flagged ambiguous cells.')
